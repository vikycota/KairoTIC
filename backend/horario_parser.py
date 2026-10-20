"""Lectura y normalización de archivos de bloques horarios."""
import csv
import io
import json
import re
import unicodedata
from datetime import datetime, time

from openpyxl import Workbook, load_workbook


class HorarioError(Exception):
    def __init__(self, message):
        self.message = message
        super().__init__(message)


ALIAS = {
    "materia": {"materia", "nombre", "asignatura", "curso"},
    "anio_carrera": {"aniocarrera", "anocarrera", "aniodelacarrera", "anodelacarrera", "anio", "ano"},
    "grupo": {"grupo", "comision", "comisiongrupo"},
    "dia": {"dia", "diadelasemana"},
    "hora_inicio": {"horainicio", "inicio", "desde"},
    "hora_fin": {"horafin", "fin", "hasta"},
    "docente": {"docente", "profesor", "profesores"},
    "tipo": {"tipo", "modalidad"},
    "opcional": {"opcional", "optativa"},
}
OBLIGATORIAS = ("materia", "dia", "hora_inicio", "hora_fin")
DIAS = {"lunes": "Lunes", "martes": "Martes", "miercoles": "Miércoles",
        "jueves": "Jueves", "viernes": "Viernes", "sabado": "Sábado"}
COLUMNAS = tuple(ALIAS)


def _clave(valor):
    valor = unicodedata.normalize("NFD", str(valor or "").strip().lower())
    return re.sub(r"[^a-z0-9]", "", "".join(c for c in valor if unicodedata.category(c) != "Mn"))


def _limpiar(valor):
    return re.sub(r"\s+", " ", str(valor).replace("\x00", " ")).strip() if valor is not None else ""


def _hora(valor, fila, campo):
    try:
        if isinstance(valor, datetime):
            valor = valor.time()
        if isinstance(valor, time):
            resultado = valor.strftime("%H:%M")
        else:
            texto = _limpiar(valor)
            coincidencia = re.fullmatch(r"(\d{1,2}):(\d{2})(?::\d{2}(?:\.\d+)?)?", texto)
            if not coincidencia:
                raise ValueError
            h, m = int(coincidencia.group(1)), int(coincidencia.group(2))
            if h > 23 or m > 59:
                raise ValueError
            resultado = f"{h:02d}:{m:02d}"
        return resultado
    except (ValueError, TypeError):
        raise HorarioError(f"Fila {fila}: '{campo}' debe tener formato HH:MM.")


def _bool(valor, fila):
    if valor is None or (isinstance(valor, str) and not valor.strip()):
        return False
    if isinstance(valor, bool):
        return valor
    clave = _clave(valor)
    if clave in {"true", "1", "si", "yes"}:
        return True
    if clave in {"false", "0", "no"}:
        return False
    raise HorarioError(f"Fila {fila}: 'opcional' debe ser verdadero o falso.")


def _entero_opcional(valor, fila):
    if valor is None or (isinstance(valor, str) and not valor.strip()):
        return None
    try:
        numero = float(str(valor).strip().replace(",", "."))
        if numero.is_integer() and 1 <= numero <= 2147483647:
            return int(numero)
    except (ValueError, TypeError):
        pass
    raise HorarioError(f"Fila {fila}: 'anio_carrera' debe ser un entero positivo o quedar vacío.")


def normalizar_filas(filas_crudas):
    resultado = []
    errores = []
    for numero, cruda in filas_crudas:
        if not any(_limpiar(valor) for valor in cruda.values()):
            continue
        materia = _limpiar(cruda.get("materia"))
        dia = DIAS.get(_clave(cruda.get("dia")))
        fila = {"fila": numero, "materia": materia, "dia": dia}
        try:
            fila.update({
                "anio_carrera": _entero_opcional(cruda.get("anio_carrera"), numero),
                "grupo": _limpiar(cruda.get("grupo")) or "General",
                "hora_inicio": _hora(cruda.get("hora_inicio"), numero, "hora_inicio"),
                "hora_fin": _hora(cruda.get("hora_fin"), numero, "hora_fin"),
                "docente": _limpiar(cruda.get("docente")) or None,
                "tipo": _limpiar(cruda.get("tipo")) or "Clase",
                "opcional": _bool(cruda.get("opcional"), numero),
            })
            if not materia:
                errores.append(f"Fila {numero}: falta la materia.")
            if len(materia) > 100:
                errores.append(f"Fila {numero}: la materia no puede superar 100 caracteres.")
            if len(fila["grupo"]) > 50:
                errores.append(f"Fila {numero}: el grupo no puede superar 50 caracteres.")
            if len(fila["docente"] or "") > 250:
                errores.append(f"Fila {numero}: el docente no puede superar 250 caracteres.")
            if len(fila["tipo"]) > 50:
                errores.append(f"Fila {numero}: el tipo no puede superar 50 caracteres.")
            if not dia:
                errores.append(f"Fila {numero}: día no válido. Usá lunes a sábado.")
            if fila["hora_fin"] <= fila["hora_inicio"]:
                errores.append(f"Fila {numero}: la hora final debe ser posterior a la hora inicial.")
            resultado.append(fila)
        except HorarioError as error:
            errores.append(error.message)
    if not resultado and not errores:
        errores.append("El archivo no contiene bloques horarios.")
    return resultado, errores


def _tabla(encabezados, filas):
    indices = {}
    for i, encabezado in enumerate(encabezados):
        clave = _clave(encabezado)
        for canonico, alias in ALIAS.items():
            if clave in alias:
                indices[canonico] = i
                break
    faltan = [nombre for nombre in OBLIGATORIAS if nombre not in indices]
    if faltan:
        raise HorarioError("Falta la columna obligatoria: " + ", ".join(f"'{c}'" for c in faltan) + ".")
    crudas = []
    for numero, valores in filas:
        crudas.append((numero, {campo: valores[i] if i < len(valores) else None for campo, i in indices.items()}))
    return normalizar_filas(crudas)


def leer_archivo_horarios(filename, contenido):
    extension = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if not contenido:
        raise HorarioError("El archivo está vacío.")
    if extension == "json":
        try:
            datos = json.loads(contenido.decode("utf-8-sig"))
        except (UnicodeDecodeError, json.JSONDecodeError):
            raise HorarioError("El archivo JSON no es válido o no está codificado en UTF-8.")
        if isinstance(datos, dict):
            datos = datos.get("horarios")
        if not isinstance(datos, list):
            raise HorarioError("El JSON debe ser una lista de bloques o un objeto con la propiedad 'horarios'.")
        filas = []
        for nro, item in enumerate(datos, 1):
            if not isinstance(item, dict):
                raise HorarioError(f"Fila {nro}: cada horario debe ser un objeto JSON.")
            mapped = {}
            for key, value in item.items():
                clave = _clave(key)
                for canonico, alias in ALIAS.items():
                    if clave in alias:
                        mapped[canonico] = value
                        break
            filas.append((nro + 1, mapped))
        bloques, errores = normalizar_filas(filas)
        return bloques, errores
    if extension == "csv":
        try:
            texto = contenido.decode("utf-8-sig")
        except UnicodeDecodeError:
            raise HorarioError("El CSV debe estar codificado en UTF-8.")
        muestra = texto[:8192]
        try:
            delimitador = csv.Sniffer().sniff(muestra, delimiters=",;").delimiter
        except csv.Error:
            delimitador = ";" if muestra.splitlines()[0].count(";") > muestra.splitlines()[0].count(",") else ","
        tabla = list(csv.reader(io.StringIO(texto), delimiter=delimitador))
        inicio = next((i for i, row in enumerate(tabla) if any(_limpiar(c) for c in row)), None)
        if inicio is None:
            raise HorarioError("El archivo está vacío.")
        return _tabla(tabla[inicio], ((i + 1, row) for i, row in enumerate(tabla) if i > inicio))
    if extension == "xlsx":
        try:
            libro = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
            hoja = libro.worksheets[0]
            filas = list(hoja.iter_rows(values_only=True))
            libro.close()
        except Exception:
            raise HorarioError("No se pudo abrir el Excel. Verificá que sea un .xlsx válido.")
        inicio = next((i for i, row in enumerate(filas) if any(v is not None and str(v).strip() for v in row)), None)
        if inicio is None:
            raise HorarioError("El archivo está vacío.")
        return _tabla(filas[inicio], ((i + 1, row) for i, row in enumerate(filas) if i > inicio))
    raise HorarioError("Formato no soportado. Usá CSV, XLSX o JSON.")


def plantilla_csv():
    salida = io.StringIO(newline="")
    writer = csv.writer(salida)
    writer.writerow(COLUMNAS)
    writer.writerow(["Álgebra Lineal", 1, 1, "Lunes", "08:00", "08:50", "Docente de ejemplo", "Clase", "false"])
    return "\ufeff" + salida.getvalue()


def plantilla_json():
    return json.dumps({"horarios": [{"materia": "Álgebra Lineal", "anio_carrera": 1, "grupo": "1", "dia": "Lunes", "hora_inicio": "08:00", "hora_fin": "08:50", "docente": "Docente de ejemplo", "tipo": "Clase", "opcional": False}]}, ensure_ascii=False, indent=2)


def plantilla_xlsx():
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Horarios"
    hoja.append(COLUMNAS)
    hoja.append(["Álgebra Lineal", 1, "1", "Lunes", "08:00", "08:50", "Docente de ejemplo", "Clase", False])
    salida = io.BytesIO()
    libro.save(salida)
    libro.close()
    return salida.getvalue()
