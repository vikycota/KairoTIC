import csv
import difflib
import io
import json
import re
import unicodedata

from openpyxl import load_workbook

MAX_SEMESTRE = 20          
MAX_MATERIAS = 50
MAX_CREDITOS_MATERIA = 60
MAX_NOMBRE = 100

ORDINALES = {
    "primero": 1, "primer": 1, "segundo": 2, "tercero": 3, "tercer": 3,
    "cuarto": 4, "quinto": 5, "sexto": 6, "septimo": 7, "octavo": 8,
    "noveno": 9, "decimo": 10,
}

# nombre canónico -> nombres de columna aceptados (ya normalizados)
ALIAS_COLUMNAS = {
    "semestre": {"semestre", "sem", "periodo"},
    "nombre": {"materia", "nombre", "asignatura", "curso"},
    "creditos": {"creditos", "credito", "ects", "cred"},
    "categoria": {"categoria", "area", "tipo"},
    "previas": {"previas", "previa", "previaturas", "previatura", "requisitos", "correlativas"},
    "carrera": {"carrera", "plan"},
}
COLUMNAS_OBLIGATORIAS = ("semestre", "nombre", "creditos")


class PlanError(Exception):
    """El archivo no se pudo leer (formato roto, columnas faltantes, etc.)."""

    def __init__(self, message):
        self.message = message
        super().__init__(message)


# Utilidades

def _sin_acentos(texto):
    return "".join(
        c for c in unicodedata.normalize("NFD", texto) if unicodedata.category(c) != "Mn"
    )


def _clave(texto):
    """'Créditos ' -> 'creditos' (para comparar encabezados)."""
    return re.sub(r"[^a-z0-9]", "", _sin_acentos(str(texto)).lower())


def _limpiar(texto):
    texto = re.sub(r"[\x00-\x1f\x7f]", " ", str(texto))
    return re.sub(r"\s+", " ", texto).strip()


def _vacio(valor):
    return valor is None or (isinstance(valor, str) and not valor.strip())


def _entero(valor):
    """Devuelve int o None si no es un entero (acepta 10, '10', '10.0', 10.0)."""
    if isinstance(valor, bool) or _vacio(valor):
        return None
    try:
        numero = float(str(valor).strip().replace(",", "."))
    except ValueError:
        return None
    return int(numero) if numero.is_integer() else None


def _semestre(valor):
    """
    Devuelve (ok, semestre). Vacío / '-' / 'extra' -> (True, None): materia
    fuera de los semestres (pasantías, actividades). Acepta 3, '3', 'tercero'.
    """
    if _vacio(valor):
        return True, None
    texto = _clave(valor)
    if texto in {"extra", "extras", "sin", "sinsemestre", "na", ""}:
        return True, None
    if texto in ORDINALES:
        return True, ORDINALES[texto]
    coincide = re.fullmatch(r"(?:semestre|sem)?(\d+)o?", texto)
    numero = _entero(coincide.group(1)) if coincide else _entero(valor)
    if numero is not None and 1 <= numero <= MAX_SEMESTRE:
        return True, numero
    return False, None


def _lista_previas(valor):
    if _vacio(valor):
        return []
    if isinstance(valor, (list, tuple)):
        partes = [str(p) for p in valor]
    else:
        partes = re.split(r"[;|\n]", str(valor))
    return [p for p in (_limpiar(p) for p in partes) if p and p not in {"-", "—"}]


def _filas_desde_tabla(encabezados, filas_crudas):
    """
    encabezados: lista de textos; filas_crudas: iterable de (nro_fila, [celdas]).
    """
    indice = {}
    for pos, enc in enumerate(encabezados):
        clave = _clave(enc)
        for canonico, alias in ALIAS_COLUMNAS.items():
            if clave in alias and canonico not in indice:
                indice[canonico] = pos

    faltan = [c for c in COLUMNAS_OBLIGATORIAS if c not in indice]
    if faltan:
        raise PlanError(
            "Faltan columnas obligatorias: " + ", ".join(faltan)
            + ". Se esperan: semestre, materia, creditos (y opcionales: categoria, previas, carrera)."
        )

    filas = []
    carreras = set()
    for nro, celdas in filas_crudas:
        if all(_vacio(c) for c in celdas):
            continue

        def celda(nombre):
            pos = indice.get(nombre)
            return celdas[pos] if pos is not None and pos < len(celdas) else None

        carrera = celda("carrera")
        if not _vacio(carrera):
            carreras.add(_limpiar(carrera))
        filas.append({
            "fila": nro,
            "semestre": celda("semestre"),
            "nombre": celda("nombre"),
            "creditos": celda("creditos"),
            "categoria": celda("categoria"),
            "previas": celda("previas"),
        })

    if len(carreras) > 1:
        raise PlanError(
            "La columna 'carrera' tiene más de un valor (" + ", ".join(sorted(carreras))
            + "). Importá un plan por archivo."
        )
    return {"carrera": next(iter(carreras), None), "valor_credito": None, "filas": filas}


def _decodificar_texto(contenido):
    for codificacion in ("utf-8-sig", "cp1252"):
        try:
            return contenido.decode(codificacion)
        except UnicodeDecodeError:
            continue
    raise PlanError("No se pudo leer el texto del archivo (codificación desconocida).")


def leer_csv(contenido):
    texto = _decodificar_texto(contenido)
    primera_linea = next((l for l in texto.splitlines() if l.strip()), "")
    if not primera_linea:
        raise PlanError("El archivo CSV está vacío.")
    delimitador = ";" if primera_linea.count(";") >= primera_linea.count(",") else ","
    lector = csv.reader(io.StringIO(texto), delimiter=delimitador)
    filas = list(lector)
    # el encabezado es la primera fila no vacía
    inicio = next((i for i, f in enumerate(filas) if any(c.strip() for c in f)), None)
    if inicio is None:
        raise PlanError("El archivo CSV está vacío.")
    return _filas_desde_tabla(
        filas[inicio],
        ((i + 1, f) for i, f in enumerate(filas) if i > inicio),
    )


def leer_xlsx(contenido):
    try:
        libro = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
    except Exception:
        raise PlanError("No se pudo abrir el Excel. Verificá que sea un .xlsx válido.")
    hoja = libro.worksheets[0]
    filas = [(i, list(f)) for i, f in enumerate(hoja.iter_rows(values_only=True), start=1)]
    libro.close()
    inicio = next((k for k, (_, f) in enumerate(filas) if any(not _vacio(c) for c in f)), None)
    if inicio is None:
        raise PlanError("La primera hoja del Excel está vacía.")
    return _filas_desde_tabla(filas[inicio][1], filas[inicio + 1:])


def leer_json(contenido):
    try:
        datos = json.loads(_decodificar_texto(contenido))
    except json.JSONDecodeError as e:
        raise PlanError(f"El JSON no es válido (línea {e.lineno}, columna {e.colno}): {e.msg}.")
    if not isinstance(datos, dict):
        raise PlanError("El JSON debe ser un objeto con la clave 'materias'.")

    materias = datos.get("materias")
    if materias is None and isinstance(datos.get("semestres"), list):
        # forma anidada: {"semestres": [{"numero": 1, "materias": [...]}]}
        materias = []
        for sem in datos["semestres"]:
            if not isinstance(sem, dict):
                continue
            for m in sem.get("materias") or []:
                if isinstance(m, dict):
                    materias.append({"semestre": sem.get("numero"), **m})
    if not isinstance(materias, list):
        raise PlanError("El JSON debe tener una lista 'materias' (o 'semestres' con materias).")

    filas = []
    for i, m in enumerate(materias, start=1):
        if not isinstance(m, dict):
            filas.append({"fila": i, "semestre": None, "nombre": None, "creditos": None,
                          "categoria": None, "previas": None})
            continue
        filas.append({
            "fila": i,
            "semestre": m.get("semestre"),
            "nombre": m.get("nombre", m.get("materia")),
            "creditos": m.get("creditos", m.get("ects")),
            "categoria": m.get("categoria"),
            "previas": m.get("previas"),
        })
    valor = datos.get("valor_credito")
    return {"carrera": datos.get("carrera"), "valor_credito": valor, "filas": filas}


def leer_archivo(nombre_archivo, contenido):
    """Elige el lector según la extensión."""
    extension = (nombre_archivo or "").rsplit(".", 1)[-1].lower() if "." in (nombre_archivo or "") else ""
    if extension == "csv":
        return leer_csv(contenido)
    if extension == "xlsx":
        return leer_xlsx(contenido)
    if extension == "json":
        return leer_json(contenido)
    raise PlanError("Formato no soportado. Subí un archivo .csv, .xlsx o .json.")


# ---------------------------------------------------------------------------
# Validación
# ---------------------------------------------------------------------------

def _renombrar_repetidas(materias, avisos):
    """
    La base identifica cada materia por su nombre, así que no puede haber dos
    iguales. Las repetidas (p. ej. 'Taller' en varios semestres) se renombran
    y se avisa. Devuelve el conjunto de nombres originales que quedaron ambiguos.
    """
    grupos = {}
    for m in materias:
        grupos.setdefault(m["nombre"].lower(), []).append(m)

    ambiguos = set()
    for clave, grupo in grupos.items():
        if len(grupo) < 2:
            continue
        ambiguos.add(clave)
        original = grupo[0]["nombre"]
        semestres = [m["semestre"] for m in grupo]
        por_semestre = len(set(semestres)) == len(semestres)
        for i, m in enumerate(grupo, start=1):
            if por_semestre:
                sufijo = f"Sem. {m['semestre']}" if m["semestre"] else "extra"
            else:
                sufijo = str(i)
            m["nombre"] = f"{original} ({sufijo})"[:MAX_NOMBRE]
        avisos.append(
            f"«{original}» aparece {len(grupo)} veces; cada materia necesita un nombre único, "
            "así que se renombraron: " + ", ".join(f"«{m['nombre']}»" for m in grupo) + "."
        )
    return ambiguos


def _detectar_ciclo(grafo):
    """grafo: {materia: [previas]}. Devuelve la lista del ciclo o None."""
    estado = {}   # 1 = en curso, 2 = terminado
    pila = []

    def visitar(nodo):
        estado[nodo] = 1
        pila.append(nodo)
        for previa in grafo.get(nodo, []):
            if previa not in grafo:
                continue
            if estado.get(previa) == 1:
                return pila[pila.index(previa):] + [previa]
            if estado.get(previa) is None:
                ciclo = visitar(previa)
                if ciclo:
                    return ciclo
        pila.pop()
        estado[nodo] = 2
        return None

    for nodo in grafo:
        if estado.get(nodo) is None:
            ciclo = visitar(nodo)
            if ciclo:
                return ciclo
    return None


def validar_plan(raw, carrera=None, existentes=None):
    """
    raw: resultado de leer_archivo. carrera: nombre forzado desde el formulario.
    existentes: {nombre_materia: creditos} de lo que ya hay en la base.
    """
    errores, avisos = [], []
    existentes = existentes or {}
    existentes_lower = {n.lower(): n for n in existentes}

    carrera = _limpiar(carrera or raw.get("carrera") or "")
    if not carrera:
        errores.append("Falta el nombre de la carrera (columna 'carrera', clave 'carrera' del JSON o campo del formulario).")
    elif len(carrera) > 100:
        errores.append("El nombre de la carrera no puede superar los 100 caracteres.")

    valor_credito = None
    if not _vacio(raw.get("valor_credito")):
        try:
            valor_credito = float(str(raw["valor_credito"]).replace(",", "."))
            if valor_credito < 0 or valor_credito > 99999999:
                raise ValueError
        except ValueError:
            errores.append("'valor_credito' debe ser un número positivo.")
            valor_credito = None

    filas = raw.get("filas") or []
    if not filas:
        errores.append("El archivo no tiene materias.")
    if len(filas) > MAX_MATERIAS:
        errores.append(f"El plan supera el máximo de {MAX_MATERIAS} materias.")
        filas = filas[:MAX_MATERIAS]

    # --- 1) campos de cada fila
    materias = []
    for f in filas:
        ref = f"Fila {f['fila']}"
        nombre = _limpiar(f["nombre"]) if not _vacio(f["nombre"]) else ""
        if not nombre:
            errores.append(f"{ref}: falta el nombre de la materia.")
            continue
        if len(nombre) > MAX_NOMBRE:
            errores.append(f"{ref}: el nombre «{nombre[:30]}…» supera los {MAX_NOMBRE} caracteres.")
            continue

        creditos = _entero(f["creditos"])
        if creditos is None or not (0 <= creditos <= MAX_CREDITOS_MATERIA):
            errores.append(f"{ref} («{nombre}»): los créditos deben ser un número entero entre 0 y {MAX_CREDITOS_MATERIA} (se leyó «{f['creditos']}»).")
            continue

        ok, semestre = _semestre(f["semestre"])
        if not ok:
            errores.append(f"{ref} («{nombre}»): el semestre debe ser un número de 1 a {MAX_SEMESTRE} o quedar vacío (se leyó «{f['semestre']}»).")
            continue

        categoria = _limpiar(f["categoria"]) if not _vacio(f["categoria"]) else None
        if categoria and len(categoria) > 100:
            errores.append(f"{ref} («{nombre}»): la categoría supera los 100 caracteres.")
            continue

        materias.append({
            "fila": f["fila"], "nombre": nombre, "creditos": creditos,
            "semestre": semestre, "categoria": categoria,
            "previas": _lista_previas(f["previas"]),
        })

    # --- 2) nombres repetidos
    ambiguos = _renombrar_repetidas(materias, avisos)

    # --- 3) previas
    en_plan = {m["nombre"].lower(): m for m in materias}
    for m in materias:
        resueltas = []
        for previa in m["previas"]:
            ref = f"Fila {m['fila']} («{m['nombre']}»)"
            clave = previa.lower()
            if clave in ambiguos:
                errores.append(f"{ref}: la previa «{previa}» es ambigua porque esa materia aparece varias veces en el plan.")
            elif clave == m["nombre"].lower():
                errores.append(f"{ref}: una materia no puede ser previa de sí misma.")
            elif clave in en_plan:
                destino = en_plan[clave]
                if m["semestre"] and destino["semestre"] and destino["semestre"] >= m["semestre"]:
                    avisos.append(f"{ref}: su previa «{destino['nombre']}» está en el semestre {destino['semestre']}, que no es anterior al {m['semestre']}.")
                if destino["nombre"] not in resueltas:
                    resueltas.append(destino["nombre"])
            elif clave in existentes_lower:
                if existentes_lower[clave] not in resueltas:
                    resueltas.append(existentes_lower[clave])
            else:
                candidatos = difflib.get_close_matches(
                    previa, [x["nombre"] for x in materias] + list(existentes), n=1, cutoff=0.6
                )
                pista = f" ¿Quisiste decir «{candidatos[0]}»?" if candidatos else ""
                if "," in previa:
                    pista += " Si son varias previas, separalas con punto y coma (;)."
                errores.append(f"{ref}: la previa «{previa}» no existe en el plan ni en la base.{pista}")
        m["previas"] = resueltas

    ciclo = _detectar_ciclo({m["nombre"]: m["previas"] for m in materias})
    if ciclo:
        errores.append("Hay un ciclo de previas: " + " → ".join(ciclo) + ".")

    # --- 4) choques con lo que ya está en la base (otras carreras / importaciones previas)
    for m in materias:
        actuales = existentes.get(m["nombre"])
        if actuales is None and m["nombre"].lower() in existentes_lower:
            actuales = existentes[existentes_lower[m["nombre"].lower()]]
        if actuales is not None and actuales != m["creditos"]:
            avisos.append(f"«{m['nombre']}» ya existe con {actuales} créditos; se actualizará a {m['creditos']}.")

    resumen = {
        "materias": len(materias),
        "creditos_total": sum(m["creditos"] for m in materias),
        "semestres": len({m["semestre"] for m in materias if m["semestre"]}),
        "sin_semestre": sum(1 for m in materias if not m["semestre"]),
        "con_previas": sum(1 for m in materias if m["previas"]),
    }
    return {
        "valido": not errores,
        "carrera": carrera,
        "valor_credito": valor_credito,
        "materias": materias,
        "errores": errores,
        "avisos": avisos,
        "resumen": resumen,
    }


# ---------------------------------------------------------------------------
# Plantillas descargables
# ---------------------------------------------------------------------------

PLANTILLA_FILAS = [
    ("1", "Análisis Mat. I", 10, "C. Básicas comunes", ""),
    ("1", "Programación I", 6, "Informática comunes", ""),
    ("2", "Análisis Mat. II", 10, "C. Básicas comunes", "Análisis Mat. I"),
    ("2", "Diseño de Base de Datos I", 6, "Informática comunes", "Programación I"),
    ("3", "Programación II", 6, "Ingeniería aplicada", "Programación I; Diseño de Base de Datos I"),
    ("", "Pasantía", 20, "Extras", ""),
]


def plantilla_csv():
    salida = io.StringIO()
    escritor = csv.writer(salida, delimiter=";", lineterminator="\r\n")
    escritor.writerow(["semestre", "materia", "creditos", "categoria", "previas"])
    escritor.writerows(PLANTILLA_FILAS)
    # BOM para que Excel abra bien las tildes
    return "\ufeff" + salida.getvalue()


def plantilla_json():
    return json.dumps({
        "carrera": "Ingeniería en TIC",
        "valor_credito": 45.5,
        "materias": [
            {"semestre": int(s) if s else None, "nombre": n, "creditos": c, "categoria": cat,
             "previas": [p.strip() for p in pre.split(";")] if pre else []}
            for s, n, c, cat, pre in PLANTILLA_FILAS
        ],
    }, ensure_ascii=False, indent=2)
